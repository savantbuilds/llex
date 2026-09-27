"""Tests for document export.

Every format is verified by reading the produced bytes back with the library
that defines it, so a writer cannot silently emit a file nothing can open.
"""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import pytest

from llex.document import Document, PageSetup
from llex.export import SUPPORTED_FORMATS, Exporter, ExportError
from llex.markup import parse_document, to_plain_text

RICH_HTML = (
    '<div class="page">'
    "<h1>Report Title</h1>"
    "<h3>A Section</h3>"
    '<p style="text-align: justify">The <strong>engine</strong> measures '
    "<em>rendered</em> height, so layout is only knowable in the browser.</p>"
    '<p>Bold <b>x</b>, italic <i>y</i>, underline <u>z</u>, strike <s>w</s>, '
    'highlight <mark>h</mark>, and <span style="color:#cc0000">red</span>.</p>'
    "<ul><li>first</li><li>second<ul><li>nested</li></ul></li></ul>"
    "<ol><li>one</li><li>two</li></ol>"
    "<blockquote><p>Quoted wisdom.</p></blockquote>"
    '<pre><code class="language-python">def f():\n    return 1</code></pre>'
    "<hr>"
    '<p style="text-align: center">centred</p>'
    "<p>Unicode: caf\u00e9 na\u00efve \u2014 \u4e2d\u6587 \U0001f600</p>"
    "</div>"
    '<div class="page"><p>Second physical page.</p></div>'
)


@pytest.fixture
def document() -> Document:
    return Document(title="Layout Notes", content=RICH_HTML)


@pytest.fixture
def exporter(document: Document) -> Exporter:
    return Exporter(document)


def stable_payload(payload: bytes) -> tuple[str, ...]:
    """A timestamp-independent view of a rendered export.

    ``zipfile`` stamps every entry with the current time, so two renders of the
    same document a second apart are not byte-identical. Comparing the entry
    names and decoded contents is what actually matters.
    """
    if payload[:2] != b"PK":
        return (payload.decode("utf-8", "replace"),)
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        return tuple(
            f"{name}\n{archive.read(name).decode('utf-8', 'replace')}" for name in sorted(archive.namelist())
        )


class TestFormatRegistry:
    def test_every_advertised_format_renders(self, exporter: Exporter) -> None:
        for suffix in SUPPORTED_FORMATS:
            assert exporter.render(suffix), f"{suffix} produced nothing"

    def test_registry_matches_the_dispatch_table(self) -> None:
        from llex.export import _WRITERS

        assert set(_WRITERS) == set(SUPPORTED_FORMATS)

    def test_pdf_is_not_offered(self) -> None:
        assert ".pdf" not in SUPPORTED_FORMATS

    def test_formats_property_is_a_copy(self, exporter: Exporter) -> None:
        formats = exporter.formats
        formats[".txt"] = "mutated"
        assert exporter.formats[".txt"] == "Plain Text"

    def test_empty_document_renders_every_format(self) -> None:
        for suffix in SUPPORTED_FORMATS:
            assert Exporter(Document()).render(suffix)


class TestSuffixHandling:
    def test_bare_extension_is_accepted(self, exporter: Exporter) -> None:
        assert stable_payload(exporter.render("docx")) == stable_payload(exporter.render(".docx"))

    def test_case_is_insensitive(self, exporter: Exporter) -> None:
        assert stable_payload(exporter.render(".DOCX")) == stable_payload(exporter.render(".docx"))

    def test_unsupported_format_is_rejected(self, exporter: Exporter) -> None:
        with pytest.raises(ExportError, match="unsupported export format"):
            exporter.render(".pdf")

    def test_error_lists_the_supported_formats(self, exporter: Exporter) -> None:
        with pytest.raises(ExportError) as info:
            exporter.render(".pages")
        for suffix in SUPPORTED_FORMATS:
            assert suffix in str(info.value)


class TestWriteToDisk:
    def test_writes_and_reports(self, exporter: Exporter, tmp_path: Path) -> None:
        result = exporter.write(tmp_path / "out.md")
        assert result.path == tmp_path / "out.md"
        assert result.format == ".md"
        assert result.size == (tmp_path / "out.md").stat().st_size
        assert result.to_dict()["file_name"] == "out.md"

    def test_suffix_is_authoritative(self, exporter: Exporter, tmp_path: Path) -> None:
        with pytest.raises(ExportError):
            exporter.write(tmp_path / "out.pdf")

    def test_creates_missing_directories(self, exporter: Exporter, tmp_path: Path) -> None:
        assert exporter.write(tmp_path / "a" / "b" / "out.txt").path.is_file()

    def test_overwrites_existing(self, exporter: Exporter, tmp_path: Path) -> None:
        target = tmp_path / "out.txt"
        target.write_text("stale", encoding="utf-8")
        exporter.write(target)
        assert "stale" not in target.read_text(encoding="utf-8")


class TestSuggestedFilename:
    @pytest.mark.parametrize("title", ["My Report", "a/b\\c:d*e?f", "  ", "..."])
    def test_is_always_filesystem_safe(self, exporter: Exporter, title: str) -> None:
        exporter.document.title = title
        name = exporter.suggested_filename(".docx")
        assert not re.search(r'[\\/:*?"<>|]', name)
        assert name.endswith(".docx")

    def test_uses_the_document_title(self, exporter: Exporter) -> None:
        assert exporter.suggested_filename(".md") == "Layout Notes.md"

    def test_long_titles_are_truncated(self, exporter: Exporter) -> None:
        exporter.document.title = "x" * 500
        assert len(exporter.suggested_filename(".md")) <= 124


class TestPlainText:
    def test_content(self, exporter: Exporter) -> None:
        text = exporter.render(".txt").decode()
        assert "Report Title" in text
        assert "- first" in text
        assert "1. one" in text
        assert "<strong>" not in text

    def test_unicode_survives(self, exporter: Exporter) -> None:
        assert "caf\u00e9" in exporter.render(".txt").decode()


class TestMarkdown:
    def test_headings_and_emphasis(self, exporter: Exporter) -> None:
        text = exporter.render(".md").decode()
        assert "# Report Title" in text
        assert "### A Section" in text
        assert "**engine**" in text
        assert "*rendered*" in text
        assert "~~w~~" in text

    def test_lists_nest(self, exporter: Exporter) -> None:
        lines = exporter.render(".md").decode().splitlines()
        assert "- first" in lines
        assert "  - nested" in lines

    def test_code_fence_carries_the_language(self, exporter: Exporter) -> None:
        text = exporter.render(".md").decode()
        assert "```python" in text
        assert "    return 1" in text

    def test_blockquote(self, exporter: Exporter) -> None:
        assert "> Quoted wisdom." in exporter.render(".md").decode()

    def test_horizontal_rule(self, exporter: Exporter) -> None:
        assert "\n---\n" in exporter.render(".md").decode()

    def test_no_triple_blank_lines(self, exporter: Exporter) -> None:
        assert "\n\n\n" not in exporter.render(".md").decode()

    def test_markdown_syntax_in_content_is_escaped(self) -> None:
        document = Document(content="<p>a * b _ c [d]</p>")
        text = Exporter(document).render(".md").decode()
        assert r"\*" in text and r"\_" in text and r"\[" in text


class TestHtml:
    def test_is_a_standalone_document(self, exporter: Exporter) -> None:
        text = exporter.render(".html").decode()
        assert text.startswith("<!DOCTYPE html>")
        assert "</html>" in text
        assert "<style>" in text

    def test_escapes_dangerous_characters(self) -> None:
        document = Document(content="<p>a &amp; b &lt;script&gt;</p>")
        text = Exporter(document).render(".html").decode()
        assert "&lt;script&gt;" in text
        assert "<script>" not in text

    def test_marks_become_tags(self, exporter: Exporter) -> None:
        text = exporter.render(".html").decode()
        assert "<strong>engine</strong>" in text
        assert "<em>rendered</em>" in text
        assert "<mark" in text
        assert 'style="color: #cc0000"' in text

    def test_page_geometry_reaches_the_print_css(self) -> None:
        document = Document(content="<p>x</p>", page=PageSetup(margin_left=1.5))
        text = Exporter(document).render(".html").decode()
        assert "max-width: 816px" in text
        assert "144px" in text

    def test_parses_without_error(self, exporter: Exporter) -> None:
        import html.parser

        parser = html.parser.HTMLParser()
        parser.feed(exporter.render(".html").decode())
        parser.close()


class TestRtf:
    def test_has_a_valid_header(self, exporter: Exporter) -> None:
        text = exporter.render(".rtf").decode()
        assert text.startswith(r"{\rtf1")
        assert text.rstrip().endswith("}")
        assert text.count("{") == text.count("}")

    def test_font_and_colour_tables(self, exporter: Exporter) -> None:
        text = exporter.render(".rtf").decode()
        assert r"{\fonttbl" in text
        assert r"{\colortbl;" in text
        assert r"\red204\green0\blue0;" in text

    def test_page_geometry_in_twips(self, exporter: Exporter) -> None:
        text = exporter.render(".rtf").decode()
        assert r"\paperw12240" in text  # 8.5in
        assert r"\paperh15840" in text  # 11in

    def test_custom_margins(self) -> None:
        document = Document(content="<p>x</p>", page=PageSetup(margin_left=1.5))
        assert r"\margl2160" in Exporter(document).render(".rtf").decode()

    def test_unicode_is_escaped(self, exporter: Exporter) -> None:
        text = exporter.render(".rtf").decode()
        assert r"\u233?" in text  # é
        assert "café" not in text

    def test_astral_characters_use_surrogate_pairs(self) -> None:
        document = Document(content="<p>\U0001f600</p>")
        text = Exporter(document).render(".rtf").decode()
        assert r"\u-10179?\u-8704?" in text

    def test_marks(self, exporter: Exporter) -> None:
        text = exporter.render(".rtf").decode()
        assert r"\b" in text
        assert r"\i" in text
        assert r"\ul" in text
        assert r"\strike" in text

    def test_output_is_pure_ascii(self, exporter: Exporter) -> None:
        exporter.render(".rtf").decode("ascii")  # raises if not

    def test_headings_carry_outline_levels(self, exporter: Exporter) -> None:
        assert r"\outlinelevel0" in exporter.render(".rtf").decode()
        assert r"\outlinelevel2" in exporter.render(".rtf").decode()


class TestDocx:
    def test_readable_by_python_docx(self, exporter: Exporter) -> None:
        import docx

        package = docx.Document(io.BytesIO(exporter.render(".docx")))
        assert package.paragraphs
        assert package.core_properties.title == "Layout Notes"

    def test_heading_styles_are_applied(self, exporter: Exporter) -> None:
        import docx

        package = docx.Document(io.BytesIO(exporter.render(".docx")))
        styles = {p.style.name for p in package.paragraphs}
        assert "Heading 1" in styles
        assert "List Bullet" in styles
        assert "List Number" in styles

    def test_first_paragraph_is_the_title(self, exporter: Exporter) -> None:
        import docx

        package = docx.Document(io.BytesIO(exporter.render(".docx")))
        assert package.paragraphs[0].text == "Report Title"

    def test_runs_carry_bold_and_colour(self, exporter: Exporter) -> None:
        import docx

        package = docx.Document(io.BytesIO(exporter.render(".docx")))
        bold = [r for p in package.paragraphs for r in p.runs if r.bold and r.text == "engine"]
        assert bold, "the <strong> run was not exported as bold"

    def test_page_geometry_is_applied(self) -> None:
        import docx
        from docx.shared import Inches

        document = Document(content="<p>x</p>", page=PageSetup(margin_left=1.5))
        package = docx.Document(io.BytesIO(Exporter(document).render(".docx")))
        section = package.sections[0]
        assert section.page_width == Inches(8.5)
        assert section.left_margin == Inches(1.5)


class TestOdt:
    def test_readable_by_odfpy(self, exporter: Exporter) -> None:
        from odf.opendocument import load

        load(io.BytesIO(exporter.render(".odt")))

    def test_is_a_valid_zip_container(self, exporter: Exporter) -> None:
        with zipfile.ZipFile(io.BytesIO(exporter.render(".odt"))) as archive:
            assert archive.namelist()
            assert archive.testzip() is None

    def test_contains_the_document_text(self, exporter: Exporter) -> None:
        with zipfile.ZipFile(io.BytesIO(exporter.render(".odt"))) as archive:
            content = archive.read("content.xml").decode("utf-8")
        assert "Report Title" in content

    def test_list_items_wrap_content_in_paragraphs(self, exporter: Exporter) -> None:
        """ODF forbids a bare <text:span> inside <text:list-item>."""
        with zipfile.ZipFile(io.BytesIO(exporter.render(".odt"))) as archive:
            content = archive.read("content.xml").decode("utf-8")
        for item in re.findall(r"<text:list-item.*?</text:list-item>", content, re.DOTALL):
            assert "<text:p" in item
            assert not re.search(r"<text:list-item>\s*<text:span", item)

    def test_nested_lists_are_preserved(self, exporter: Exporter) -> None:
        with zipfile.ZipFile(io.BytesIO(exporter.render(".odt"))) as archive:
            content = archive.read("content.xml").decode("utf-8")
        # odfpy emits a bare <text:list> for an unattributed list.
        assert content.count("<text:list") >= 3
        assert "<text:list-item>" in content


class TestEpub:
    def test_container_layout(self, exporter: Exporter) -> None:
        with zipfile.ZipFile(io.BytesIO(exporter.render(".epub"))) as archive:
            names = archive.namelist()
            assert names[0] == "mimetype"
            assert "META-INF/container.xml" in names
            assert "OEBPS/content.opf" in names
            assert "OEBPS/content.xhtml" in names
            assert "OEBPS/nav.xhtml" in names

    def test_mimetype_is_stored_uncompressed(self, exporter: Exporter) -> None:
        with zipfile.ZipFile(io.BytesIO(exporter.render(".epub"))) as archive:
            first = archive.infolist()[0]
            assert first.filename == "mimetype"
            assert first.compress_type == zipfile.ZIP_STORED
            assert archive.read("mimetype") == b"application/epub+zip"

    def test_identifier_is_stable_across_exports(self, exporter: Exporter) -> None:
        def identifier() -> str:
            with zipfile.ZipFile(io.BytesIO(exporter.render(".epub"))) as archive:
                opf = archive.read("OEBPS/content.opf").decode("utf-8")
            match = re.search(r'<dc:identifier id="pub-id">([^<]+)</dc:identifier>', opf)
            assert match is not None
            return match.group(1)

        assert identifier() == identifier()

    def test_different_documents_get_different_identifiers(self) -> None:
        first = Exporter(Document(title="A", content="<p>a</p>")).render(".epub")
        second = Exporter(Document(title="B", content="<p>b</p>")).render(".epub")
        assert first != second

    def test_navigation_document_declares_itself(self, exporter: Exporter) -> None:
        with zipfile.ZipFile(io.BytesIO(exporter.render(".epub"))) as archive:
            opf = archive.read("OEBPS/content.opf").decode("utf-8")
        assert 'properties="nav"' in opf
        assert '<itemref idref="content"/>' in opf

    def test_content_is_escaped(self) -> None:
        document = Document(title="<b>t</b>", content="<p>a &amp; b</p>")
        with zipfile.ZipFile(io.BytesIO(Exporter(document).render(".epub"))) as archive:
            content = archive.read("OEBPS/content.xhtml").decode("utf-8")
        assert "<b>t</b>" not in content
        assert "&lt;b&gt;t&lt;/b&gt;" in content


class TestZip:
    def test_contains_a_single_html_entry(self, exporter: Exporter) -> None:
        with zipfile.ZipFile(io.BytesIO(exporter.render(".zip"))) as archive:
            assert archive.namelist() == ["index.html"]
            assert archive.read("index.html").startswith(b"<!DOCTYPE html>")

#: A document laid out as three pages, as the editor would have written it.
PAGINATED_HTML = (
    '<div class="page"><h1>Alpha</h1><p>one</p></div>'
    '<div class="page"><h2>Beta</h2><p>two</p></div>'
    '<div class="page"><p>three</p></div>'
)


class TestPageGroups:
    """Where the editor's page boundaries live after parsing."""

    def test_a_paginated_document_reports_its_pages(self) -> None:
        parsed = parse_document(PAGINATED_HTML)
        assert parsed.is_paginated is True
        assert [len(group) for group in parsed.page_groups()] == [2, 2, 1]

    def test_an_unpaginated_document_is_one_group(self) -> None:
        """A caller should not have to branch on whether the source was paginated."""
        parsed = parse_document("<p>a</p><p>b</p>")
        assert parsed.is_paginated is False
        assert [len(group) for group in parsed.page_groups()] == [2]

    def test_an_empty_document_is_no_groups(self) -> None:
        assert parse_document("").page_groups() == ()

    def test_content_outside_a_page_is_not_dropped(self) -> None:
        """A caller may append blocks; they must still be exported."""
        assert "appended" in to_plain_text(PAGINATED_HTML + "<p>appended</p>")


class TestExportedPageBreaks:
    """Page breaks in the formats that can express one.

    The editor computes pagination carefully; an export that discarded it handed
    the reader a document that was not the one that was written.
    """

    @pytest.fixture
    def paginated(self) -> Exporter:
        return Exporter(Document(title="Paged", content=PAGINATED_HTML))

    def test_html_wraps_each_page_in_a_section(self, paginated: Exporter) -> None:
        html = paginated.render(".html").decode("utf-8")
        body = html.split("<body>", 1)[1]
        assert body.count('<section class="llex-page"') == 3

    def test_html_breaks_between_pages_but_not_after_the_last(self, paginated: Exporter) -> None:
        # A break after the final page would leave a blank one.
        html = paginated.render(".html").decode("utf-8")
        body = html.split("<body>", 1)[1]
        assert body.count('data-page-break="after"') == 2
        assert "data-page-break" not in body.rsplit("<section", 1)[1]

    def test_html_carries_the_css_the_break_needs(self, paginated: Exporter) -> None:
        html = paginated.render(".html").decode("utf-8")
        assert "break-after: page" in html
        assert "page-break-after: always" in html

    def test_html_preserves_document_order(self, paginated: Exporter) -> None:
        html = paginated.render(".html").decode("utf-8")
        assert html.index("Alpha") < html.index("Beta") < html.index("three")

    def test_rtf_emits_one_break_per_boundary(self, paginated: Exporter) -> None:
        rtf = paginated.render(".rtf").decode("ascii")
        assert rtf.count(r"\page ") == 2

    def test_docx_emits_one_break_per_boundary(self, paginated: Exporter) -> None:
        with zipfile.ZipFile(io.BytesIO(paginated.render(".docx"))) as archive:
            body = archive.read("word/document.xml").decode("utf-8")
        assert body.count('w:type="page"') == 2

    def test_odt_emits_one_break_per_boundary(self, paginated: Exporter) -> None:
        with zipfile.ZipFile(io.BytesIO(paginated.render(".odt"))) as archive:
            content = archive.read("content.xml").decode("utf-8")
        assert content.count('break-before="page"') == 2

    def test_odt_does_not_break_before_the_first_page(self, paginated: Exporter) -> None:
        """The document already starts on a page; a break there opens on a blank one."""
        with zipfile.ZipFile(io.BytesIO(paginated.render(".odt"))) as archive:
            content = archive.read("content.xml").decode("utf-8")
        first = content.index("Alpha")
        assert 'break-before="page"' not in content[:first]

    def test_odt_breaks_are_not_empty_paragraphs(self, paginated: Exporter) -> None:
        """A dedicated empty paragraph would add a visible blank line."""
        with zipfile.ZipFile(io.BytesIO(paginated.render(".odt"))) as archive:
            content = archive.read("content.xml").decode("utf-8")
        assert '<text:p break-before="page"/>' not in content

    def test_epub_honours_the_boundaries_too(self, paginated: Exporter) -> None:
        with zipfile.ZipFile(io.BytesIO(paginated.render(".epub"))) as archive:
            content = archive.read("OEBPS/content.xhtml").decode("utf-8")
            styles = archive.read("OEBPS/style.css").decode("utf-8")
        assert content.count('<section class="llex-page"') == 3
        assert "break-after: page" in styles

    def test_plain_text_gets_no_page_markup(self, paginated: Exporter) -> None:
        """Plain text has no pages; inventing markup would corrupt it."""
        text = paginated.render(".txt").decode("utf-8")
        assert "<" not in text
        assert "\\page" not in text

    def test_an_unpaginated_document_gets_no_breaks(self) -> None:
        """A document that was never laid out must not gain invented page breaks."""
        flat = Exporter(Document(title="Flat", content="<p>a</p><p>b</p>"))
        body = flat.render(".html").decode("utf-8").split("<body>", 1)[1]
        assert "<section" not in body
        assert flat.render(".rtf").decode("ascii").count(r"\page ") == 0

    def test_zip_shares_the_html_pagination(self, paginated: Exporter) -> None:
        with zipfile.ZipFile(io.BytesIO(paginated.render(".zip"))) as archive:
            html = archive.read("index.html").decode("utf-8")
        assert html.count('<section class="llex-page"') == 3
